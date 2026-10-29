"""Render the report's primary Excel sheets in the portal and HTML export."""
from datetime import date, datetime
from html import escape
from urllib.parse import urlsplit

from openpyxl.utils import get_column_letter


WORKBOOK_STYLE = """
.workbook-report{margin:28px 0}.workbook-tabs{display:flex;gap:8px;flex-wrap:wrap;margin:16px 0}
.workbook-tabs a{padding:8px 13px;border:1px solid #cbd5e1;border-radius:6px;background:white;text-decoration:none;font-weight:600}
.workbook-sheet{margin:20px 0;scroll-margin-top:20px}.workbook-sheet summary{cursor:pointer;font-size:18px;font-weight:650;padding:12px 0}
.workbook-sheet .scroll{border:1px solid #cbd5e1;border-radius:8px;max-height:620px;overflow:auto}
.workbook-sheet table{width:max-content;min-width:100%;border-collapse:separate;border-spacing:0;background:white;font-size:13px}
.workbook-sheet td,.workbook-sheet th{box-sizing:border-box;padding:8px;min-width:76px;max-width:340px;border-right:1px solid #d8dee8;border-bottom:1px solid #d8dee8;text-align:center;white-space:pre-wrap;overflow-wrap:anywhere;color:#172b4d;vertical-align:middle}
.workbook-sheet th{background:#edf2f7;font-weight:600}.workbook-sheet thead{position:sticky;top:0;z-index:2}
.workbook-sheet td{height:35px}.workbook-sheet td:first-child{white-space:nowrap}
.workbook-sheet a{color:#0369a1}.workbook-note{color:#607086;font-size:13px}
"""


def _color(color):
    if color is not None and color.type == 'rgb':
        value = color.rgb[-6:]
        if len(value) == 6 and all(char in '0123456789abcdefABCDEF' for char in value):
            return '#' + value
    return None


def workbook_preview_html(book):
    """Use the exported cells and merged headers so the web view matches Excel."""
    names = [name for name in book.sheetnames
             if name not in {'Summary', 'Daily Summary', 'Runs', 'Results', 'URLs', 'Errors', 'Notes'}]
    first_publication = next((name for name in names if name not in {'BenchMark', 'URL', "URL's"}), None)
    sections, links = [], []
    for index, name in enumerate(names):
        sheet = book[name]
        anchor = 'workbook-sheet-' + str(index)
        links.append(f'<a href="#{anchor}" onclick="document.getElementById(\'{anchor}\').open=true">{escape(name)}</a>')
        header_rows = 0 if name == 'BenchMark' else 1 if name in {'URL', "URL's"} else 3
        columns = min(sheet.max_column, 10 if name == 'BenchMark' else sheet.max_column if name in {'URL', "URL's"} else 29)
        merged, covered = {}, set()
        for span in sheet.merged_cells.ranges:
            if span.max_col > columns:
                continue
            merged[(span.min_row, span.min_col)] = (span.max_row - span.min_row + 1, span.max_col - span.min_col + 1)
            covered.update((row, col) for row in range(span.min_row, span.max_row + 1)
                           for col in range(span.min_col, span.max_col + 1)
                           if (row, col) != (span.min_row, span.min_col))
        widths = []
        for col in range(1, columns + 1):
            dimension = sheet.column_dimensions.get(get_column_letter(col))
            width = dimension.width if dimension else sheet.sheet_format.defaultColWidth or 13
            widths.append(f'<col style="width:{max(80, min(340, round(width * 7 + 8)))}px">')
        rows = []
        for row in range(1, sheet.max_row + 1):
            cells = []
            for col in range(1, columns + 1):
                if (row, col) in covered:
                    continue
                cell = sheet.cell(row, col)
                value = cell.value
                if isinstance(value, (datetime, date)):
                    label = value.strftime('%d-%b-%Y')
                elif isinstance(value, float):
                    label = f'{value:.3f}'.rstrip('0').rstrip('.')
                else:
                    label = '' if value is None else str(value)
                content = escape(label) or '&nbsp;'
                if cell.hyperlink:
                    target = cell.hyperlink.target or ''
                    if urlsplit(target).scheme in {'http', 'https'}:
                        content = f'<a href="{escape(target, quote=True)}">{content}</a>'
                styles = []
                background = _color(cell.fill.fgColor) if cell.fill.patternType == 'solid' else None
                foreground = _color(cell.font.color)
                if background:
                    styles.append('background:' + background)
                if foreground:
                    styles.append('color:' + foreground)
                if cell.font.bold:
                    styles.append('font-weight:700')
                if cell.alignment.horizontal in {'left', 'center', 'right'}:
                    styles.append('text-align:' + cell.alignment.horizontal)
                height, width = merged.get((row, col), (1, 1))
                attributes = f' rowspan="{height}" colspan="{width}"'
                if styles:
                    attributes += f' style="{";".join(styles)}"'
                if cell.comment:
                    attributes += f' title="{escape(cell.comment.text, quote=True)}"'
                tag = 'th' if row <= header_rows else 'td'
                if tag == 'th':
                    attributes += ' scope="colgroup"' if width > 1 else ' scope="col"'
                cells.append(f'<{tag}{attributes}>{content}</{tag}>')
            rows.append('<tr>' + ''.join(cells) + '</tr>')
        table = '<table><colgroup>' + ''.join(widths) + '</colgroup>'
        if header_rows:
            table += '<thead>' + ''.join(rows[:header_rows]) + '</thead>'
        table += '<tbody>' + ''.join(rows[header_rows:]) + '</tbody></table>'
        opened = ' open' if name == first_publication else ''
        sections.append(f'<details class="workbook-sheet" id="{anchor}"{opened}>'
                        f'<summary>{escape(name)}</summary><div class="scroll">{table}</div></details>')
    return ('<section class="workbook-report"><nav class="workbook-tabs" aria-label="Report sheets">'
            + ''.join(links) + '</nav>' + ''.join(sections) + '</section>')
