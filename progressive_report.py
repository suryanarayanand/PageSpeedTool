"""Keep a durable history and rebuild one cumulative Core Web Vitals report."""
from collections import defaultdict
from contextlib import contextmanager
from copy import copy
from datetime import datetime, timedelta, timezone
from html import escape
import json
import os
from pathlib import Path
import re
import tempfile

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from excel_layout import TEMPLATE, reference_workbook
from report_preview import workbook_preview_html
from storage import DATA_ROOT, OUTPUT

TERMINAL = {'Complete', 'Partial', 'Failed', 'Stopped', 'Interrupted'}
INDIA = timezone(timedelta(hours=5, minutes=30))


def history_root(output):
    output = Path(output)
    return DATA_ROOT / 'history' if output.resolve() == OUTPUT.resolve() else output / '_history'


def run_date(state):
    started = datetime.fromisoformat(state['started'])
    if started.tzinfo is None:
        started = started.replace(tzinfo=timezone.utc)
    return started.astimezone(INDIA).date()


@contextmanager
def history_lock(folder):
    folder.mkdir(parents=True, exist_ok=True)
    with (folder / '.lock').open('a+b') as handle:
        if os.name == 'nt':
            import msvcrt
            if handle.tell() == 0:
                handle.write(b'0')
                handle.flush()
            handle.seek(0)
            msvcrt.locking(handle.fileno(), msvcrt.LK_LOCK, 1)
        else:
            import fcntl
            fcntl.flock(handle.fileno(), fcntl.LOCK_EX)
        try:
            yield
        finally:
            if os.name == 'nt':
                handle.seek(0)
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


def _archive(state, folder):
    if state.get('status') not in TERMINAL:
        return
    identifier = state['id']
    if not re.fullmatch(r'[A-Za-z0-9_-]{1,100}', identifier):
        raise ValueError('Invalid history run ID')
    run_date(state)
    archive = folder / 'runs'
    archive.mkdir(exist_ok=True)
    path = archive / (identifier + '.json')
    # Updating the same run (for example after resume) replaces its snapshot.
    # A different run always retains its own dated history entry.
    content = json.dumps(state, indent=2, ensure_ascii=False)
    if path.exists() and path.read_text(encoding='utf-8') == content:
        return
    temporary = path.with_suffix('.tmp')
    temporary.write_text(content, encoding='utf-8')
    temporary.replace(path)


def cumulative_workbook(states):
    from engine import METRICS, FIELD, safe_cell, table_sheet

    book = load_workbook(TEMPLATE)
    publications = ['TH', 'BL', 'FL', 'SS']
    for state in states:
        for site in state['config']['feeds']:
            if site not in publications:
                book.copy_worksheet(book['TH']).title = site
                publications.append(site)
    urls = book['URL']
    urls.column_dimensions['A'].width = 14
    urls.freeze_panes = 'B2'
    for col, site in enumerate(publications, 2):
        urls.cell(1, col, site)
        urls.column_dimensions[get_column_letter(col)].width = 48
    for site in publications:
        sheet = book[site]
        sheet.freeze_panes = 'B4'
        sheet.column_dimensions['A'].width = 13
        sheet.print_title_rows = '1:3'
        sheet.sheet_view.zoomScale = 80

    offset = 0
    run_rows, result_rows, error_rows = [], [], []
    daily = defaultdict(list)
    for state in states:
        day = run_date(state)
        source = reference_workbook(state, METRICS, FIELD, safe_cell)
        source_url_columns = {cell.value: cell.column for cell in source['URL'][1]}
        samples = {site: [sample for sample in state['samples'] if sample['Publication'] == site]
                   for site in publications}
        length = max(1, max((len(rows) for rows in samples.values()), default=0))
        for index in range(length):
            excel_row, url_row = offset + index + 4, offset + index + 2
            urls.cell(url_row, 1, day).number_format = 'd-mmm-yy'
            urls.row_dimensions[url_row].height = 48
            for col, site in enumerate(publications, 2):
                sheet = book[site]
                sheet.row_dimensions[excel_row].height = 22
                if index < len(samples[site]):
                    for column in range(1, 30):
                        old = source[site].cell(index + 4, column)
                        new = sheet.cell(excel_row, column, old.value)
                        new.font = copy(old.font)
                        new.fill = copy(old.fill)
                        new.border = copy(old.border)
                        new.alignment = copy(old.alignment)
                        new.number_format = old.number_format
                        new.protection = copy(old.protection)
                        if old.comment:
                            new.comment = copy(old.comment)
                    original_url = source['URL'].cell(index + 2, source_url_columns[site])
                    target_url = urls.cell(url_row, col, original_url.value)
                    target_url.font = copy(original_url.font)
                    target_url.fill = copy(original_url.fill)
                    target_url.border = copy(original_url.border)
                    target_url.alignment = copy(original_url.alignment)
                    target_url.number_format = original_url.number_format
                    if original_url.hyperlink:
                        target_url.hyperlink = copy(original_url.hyperlink)
                    if original_url.comment:
                        target_url.comment = copy(original_url.comment)
                else:
                    reason = ('No article at this position in this run.' if samples[site] else
                              'No articles collected.' if site in state['config']['feeds'] else
                              'Publication not selected or no RSS feed configured.')
                    for column in range(2, 30):
                        cell = sheet.cell(excel_row, column, 'N/A')
                        cell.fill = PatternFill('solid', fgColor='F2F2F2')
                        cell.font = Font(name='Aptos Narrow', size=11, italic=True, color='777777')
                        cell.alignment = Alignment(horizontal='center')
                        cell.comment = Comment(reason + '\nRun: ' + state['id'], 'PageSpeed Reporter')
                date_cell = sheet.cell(excel_row, 1, day)
                date_cell.number_format = 'd-mmm-yy'
                date_cell.hyperlink = f"#'URL'!{get_column_letter(col)}{url_row}"
                detail = date_cell.comment.text + '\n' if date_cell.comment else ''
                date_cell.comment = Comment(detail + 'Run: ' + state['id'] + '\nStarted UTC: ' + state['started'], 'PageSpeed Reporter')
        offset += length
        source.close()
        results = state['results']
        run_rows.append([day, state['id'], state['status'], len(state['samples']), len(results),
                         sum(row['Status'] == 'OK' for row in results), state.get('message', '')])
        for row in results:
            daily[(day, row['Publication'], row['Device'])].append(row)
            result_rows.append([day, state['id'], row['Publication'], row['URL'], row['Device'], row['Status']]
                               + [row.get(metric) for metric in METRICS]
                               + [row.get('Field source'), row.get('PageSpeed link'), row.get('Error')])
            if row.get('Error'):
                error_rows.append([day, state['id'], row['Publication'], row['Device'] + ' / ' + row['URL'], row['Error']])
        error_rows.extend([day, state['id'], error.get('Publication'), error.get('Stage'), error.get('Error')]
                          for error in state.get('errors', []))
    for site in publications:
        book[site].print_area = f'A1:AC{max(4, offset + 3)}'
    urls.auto_filter.ref = f'A1:{get_column_letter(len(publications) + 1)}{max(2, offset + 1)}'
    summaries = []
    for (day, site, device), rows in sorted(daily.items()):
        scores = [row['Performance'] for row in rows if isinstance(row.get('Performance'), (int, float))]
        summaries.append([day, site, device, len(rows), sum(row['Status'] == 'OK' for row in rows),
                          sum(row['Status'] == 'Partial' for row in rows), sum(scores) / len(scores) if scores else None])
    table_sheet(book, 'Daily Summary', ['Date', 'Publication', 'Device', 'Checks', 'Complete', 'Partial', 'Average performance'], summaries)
    table_sheet(book, 'Runs', ['Date', 'Run ID', 'Status', 'Articles', 'Checks', 'Complete', 'Details'], run_rows)
    table_sheet(book, 'Results', ['Date', 'Run ID', 'Publication', 'URL', 'Device', 'Status'] + METRICS
                + ['Field source', 'PageSpeed link', 'Error'], result_rows)
    table_sheet(book, 'Errors', ['Date', 'Run ID', 'Publication', 'Stage', 'Error'], error_rows)
    table_sheet(book, 'Notes', ['Topic', 'Details'], [
        ['History', 'Every completed article-count run adds dated rows. Earlier runs are retained.'],
        ['Dates', 'Run start dates use India time. Multiple runs on one day remain separate.'],
        ['Missing values', 'N/A means unavailable, not zero. Failed runs do not replace earlier measurements.'],
        ['Sampling', 'Each run samples the selected article count per publication. Different days may use different articles.'],
        ['Daily averages', 'Performance averages include only measured numeric scores. Field and lab data remain separate.'],
        ['Traceability', 'The Runs, Results, and Errors sheets identify the source run.'],
    ])
    for name in ('Daily Summary', 'Runs', 'Results', 'Errors'):
        for row in book[name].iter_rows(min_row=2, max_col=1):
            row[0].number_format = 'd-mmm-yy'
    book.active = book.sheetnames.index('TH')
    return book


def _publish(folder):
    from engine import STYLE, METRICS

    states = [json.loads(path.read_text(encoding='utf-8')) for path in (folder / 'runs').glob('*.json')]
    states.sort(key=lambda state: (state['started'], state['id']))
    book = cumulative_workbook(states)
    days = sorted({run_date(state).isoformat() for state in states})
    summary = {'runs': len(states), 'days': len(days), 'first_date': days[0] if days else None,
               'latest_date': days[-1] if days else None, 'articles': sum(len(state['samples']) for state in states),
               'checks': sum(len(state['results']) for state in states),
               'measured_checks': sum(any(row.get(metric) is not None for metric in METRICS)
                                      for state in states for row in state['results']),
               'updated': datetime.now(timezone.utc).isoformat()}
    page = ('<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">'
            '<title>Daily Core Web Vitals history</title><style>' + STYLE + '</style><main>'
            '<a href="/">← Back to portal</a><h1>Daily Core Web Vitals history</h1>'
            f'<p>{summary["days"]} days · {summary["runs"]} runs · {summary["articles"]} articles · '
            f'{summary["measured_checks"]}/{summary["checks"]} checks with measurements.</p>'
            '<p>Each article-count run adds dated rows. Earlier results are retained. N/A means a measurement was unavailable. '
            'Dates use India time; the date cell identifies the source run. Different days may use different sampled articles.</p>'
            '<a class="button" href="report.xlsx">Download progressive Excel</a>'
            + workbook_preview_html(book) + '<h2>Run history</h2><div class="scroll"><table><thead><tr>'
            '<th>Date</th><th>Run</th><th>Status</th><th>Details</th></tr></thead><tbody>'
            + ''.join('<tr>' + ''.join('<td>' + escape(str(value)) + '</td>' for value in
                       (run_date(state), state['id'], state['status'], state.get('message', ''))) + '</tr>' for state in states)
            + '</tbody></table></div></main></html>')
    # Prepare all outputs before replacing any published file. JSON snapshots
    # remain the source of truth even if a later export attempt is interrupted.
    with tempfile.TemporaryDirectory(prefix='.build-', dir=folder) as temporary:
        stage = Path(temporary)
        book.save(stage / 'report.xlsx')
        book.close()
        (stage / 'report.html').write_text(page, encoding='utf-8')
        (stage / 'summary.json').write_text(json.dumps(summary, indent=2), encoding='utf-8')
        for name in ('report.xlsx', 'report.html', 'summary.json'):
            (stage / name).replace(folder / name)
    return summary


def update_progressive(state, output):
    folder = history_root(output)
    with history_lock(folder):
        _archive(state, folder)
        return _publish(folder)


def rebuild_progressive(output):
    """Backfill completed saved runs while preserving independently archived days."""
    folder = history_root(output)
    with history_lock(folder):
        for path in sorted(Path(output).glob('*/run.json')):
            state = json.loads(path.read_text(encoding='utf-8'))
            _archive(state, folder)
        return _publish(folder)
