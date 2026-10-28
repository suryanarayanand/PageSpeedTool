"""Verify a completed real report against its saved PageSpeed observations."""
import argparse
import json
import math
from pathlib import Path

from openpyxl import load_workbook
from engine import METRICS


def verify(folder):
    state = json.loads((folder / 'run.json').read_text(encoding='utf-8'))
    assert state['status'] in {'Complete', 'Partial', 'Failed'}, 'Run has not finished'
    assert len(state['results']) == len(state['samples']) * len(state['config']['strategies'])
    wb = load_workbook(folder / 'report.xlsx', data_only=True)
    sheet = wb['Results']
    headers = [c.value for c in sheet[1]]
    assert sheet.max_row == len(state['results']) + 1
    raw = [json.loads(p.read_text(encoding='utf-8')) for p in (folder / 'raw').glob('*browser.json')]
    checked = 0
    for index, result in enumerate(state['results'], 2):
        assert sheet.cell(index, headers.index('URL') + 1).value == result['URL']
        assert sheet.cell(index, headers.index('Device') + 1).value == result['Device']
        for metric in METRICS:
            expected = result.get(metric)
            actual = sheet.cell(index, headers.index(metric) + 1).value
            assert actual == 'N/A' if expected is None else isinstance(actual, (int, float)) and math.isclose(actual, expected, abs_tol=1e-8), (index, metric, expected, actual)
            checked += 1
        if result['Status'] in {'OK', 'Partial'}:
            assert any(record['PageSpeed link'] == result['PageSpeed link'] and all(record.get(k) == result.get(k) for k in METRICS) and record['Field source'] == result['Field source'] for record in raw), 'No matching captured observation'
    primary_checked = 0
    for site in state['config']['feeds']:
        samples = [s for s in state['samples'] if s['Publication'] == site]
        sheet = wb[site]
        assert sheet['B1'].value == 'Mobile' and sheet['P1'].value == 'Desktop'
        for index, sample in enumerate(samples, 4):
            for device, start in [('mobile', 2), ('desktop', 16)]:
                result = next((r for r in state['results'] if r['Publication'] == site and r['URL'] == sample['URL'] and r['Device'] == device), {})
                for offset, metric in enumerate(METRICS):
                    value = result.get(metric)
                    actual = sheet.cell(index, start + offset).value
                    assert actual == 'N/A' if value is None else isinstance(actual, (int, float)) and math.isclose(actual, value, abs_tol=1e-8)
                    primary_checked += 1
    page = (folder / 'report.html').read_text(encoding='utf-8')
    assert 'report.xlsx' in page
    assert 'No measurements collected</h1>' not in page or state['status'] == 'Failed'
    assert 'CWV passed' not in page and 'Core Web Vitals:' not in page
    print(f'Verified {len(state["results"])} device results and {checked} Excel metric cells against captured PageSpeed values.')
    print(f'Also verified {primary_checked} metric cells in the original Mobile/Desktop column positions.')
    print([(r['Publication'], r['Device'], r['Status'], r['Performance']) for r in state['results']])


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('folder', type=Path)
    verify(parser.parse_args().folder)
