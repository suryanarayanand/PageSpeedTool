import json
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

import engine
import server
from progressive_report import history_root, rebuild_progressive, update_progressive


def saved_run(identifier, started, score=90, sites=('TH',), count=1, failed=False):
    config = engine.validate_config({'feeds': {site: 'https://example.com/feed' for site in sites},
                                     'sample_size': count, 'strategies': ['mobile', 'desktop'], 'collector': 'browser'})
    samples, results = [], []
    for site in sites:
        for index in range(count):
            sample = {'Publication': site, 'URL': f'https://example.com/{site.lower()}/{index}',
                      'Title': f'Article {index + 1}', 'RSS feed': 'https://example.com/feed'}
            samples.append(sample)
            for device in config['strategies']:
                row = engine.extract_psi({})
                row.update(sample, Device=device, Status='Error' if failed else 'OK',
                           Error='PageSpeed could not resolve the article.' if failed else '',
                           Performance=None if failed else score, **{'PageSpeed link': 'https://pagespeed.web.dev/'})
                results.append(row)
    return {'id': identifier, 'started': started, 'finished': started, 'seed': 1, 'config': config,
            'status': 'Failed' if failed else 'Complete', 'message': 'No measurements collected.' if failed else 'Run complete.',
            'samples': samples, 'results': results, 'errors': []}


class ProgressiveTests(unittest.TestCase):
    def test_repeated_urls_on_different_days_keep_distinct_values_and_dates(self):
        first = saved_run('first', '2026-10-27T20:00:00+00:00', score=90)
        second = saved_run('second', '2026-10-28T20:00:00+00:00', score=42)
        with tempfile.TemporaryDirectory() as tmp:
            update_progressive(first, tmp)
            update_progressive(second, tmp)
            summary = update_progressive(first, tmp)
            self.assertEqual((summary['runs'], summary['days'], summary['articles'], summary['checks']), (2, 2, 2, 4))
            book = load_workbook(history_root(tmp) / 'report.xlsx')
            self.assertEqual(book['TH'].max_row, 5)
            self.assertEqual([book['TH'].cell(row, 7).value for row in (4, 5)], [90, 42])
            self.assertEqual([book['TH'].cell(row, 21).value for row in (4, 5)], [90, 42])
            self.assertEqual(book['TH']['A4'].value.date().isoformat(), '2026-10-28')
            self.assertEqual(book['TH']['A5'].value.date().isoformat(), '2026-10-29')
            self.assertEqual(book['TH']['A5'].hyperlink.target, "#'URL'!B3")
            self.assertEqual(book['URL']['B2'].value, book['URL']['B3'].value)
            self.assertIn('Run: second', book['TH']['A5'].comment.text)
            book.close()

    def test_failed_day_adds_na_without_replacing_successful_measurements(self):
        first = saved_run('good', '2026-10-27T09:00:00+00:00')
        failed = saved_run('failed', '2026-10-28T09:00:00+00:00', failed=True)
        with tempfile.TemporaryDirectory() as tmp:
            update_progressive(first, tmp)
            summary = update_progressive(failed, tmp)
            self.assertEqual((summary['runs'], summary['measured_checks']), (2, 2))
            book = load_workbook(history_root(tmp) / 'report.xlsx')
            self.assertEqual(book['TH']['G4'].value, 90)
            self.assertEqual(book['TH']['G5'].value, 'N/A')
            self.assertEqual(book['Runs']['C3'].value, 'Failed')
            self.assertIn('could not resolve', book['TH']['G5'].comment.text)
            self.assertEqual(book['Errors'].max_row, 3)
            book.close()

    def test_rebuild_keeps_archived_days_after_individual_run_is_removed(self):
        state = saved_run('archived', '2026-10-27T09:00:00+00:00')
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp) / 'archived'
            folder.mkdir()
            (folder / 'run.json').write_text(json.dumps(state))
            rebuild_progressive(tmp)
            shutil.rmtree(folder)
            summary = rebuild_progressive(tmp)
            self.assertEqual(summary['runs'], 1)
            self.assertTrue((history_root(tmp) / 'runs/archived.json').exists())
        self.assertEqual(history_root(engine.OUTPUT), engine.ROOT / 'history')

    def test_export_failure_preserves_previous_report_and_can_be_rebuilt(self):
        first = saved_run('first', '2026-10-27T09:00:00+00:00')
        second = saved_run('second', '2026-10-28T09:00:00+00:00')
        with tempfile.TemporaryDirectory() as tmp:
            update_progressive(first, tmp)
            folder = history_root(tmp)
            original = (folder / 'report.xlsx').read_bytes()
            with patch('openpyxl.workbook.workbook.Workbook.save', side_effect=OSError('Disk unavailable')):
                with self.assertRaises(OSError):
                    update_progressive(second, tmp)
            self.assertEqual((folder / 'report.xlsx').read_bytes(), original)
            self.assertTrue((folder / 'runs/second.json').exists())
            self.assertEqual(rebuild_progressive(tmp)['runs'], 2)

    def test_different_article_counts_keep_publication_and_url_rows_aligned(self):
        first = saved_run('first', '2026-10-27T09:00:00+00:00', sites=('TH', 'BL'), count=2)
        first['samples'] = [row for row in first['samples'] if row['URL'] != 'https://example.com/bl/1']
        first['results'] = [row for row in first['results'] if row['URL'] != 'https://example.com/bl/1']
        second = saved_run('second', '2026-10-28T09:00:00+00:00')
        with tempfile.TemporaryDirectory() as tmp:
            update_progressive(first, tmp)
            update_progressive(second, tmp)
            book = load_workbook(history_root(tmp) / 'report.xlsx')
            self.assertEqual(book['TH'].max_row, 6)
            self.assertEqual(book['URL']['B3'].value, 'https://example.com/th/1')
            self.assertIsNone(book['URL']['C3'].value)
            self.assertEqual(book['BL']['G5'].value, 'N/A')
            self.assertEqual(book['TH']['A6'].hyperlink.target, "#'URL'!B4")
            self.assertEqual(book['URL']['B4'].value, 'https://example.com/th/0')
            book.close()

    def test_custom_publications_keep_urls_when_the_feed_order_changes(self):
        first = saved_run('first', '2026-10-27T09:00:00+00:00', sites=('X', 'Y'))
        second = saved_run('second', '2026-10-28T09:00:00+00:00', sites=('Y', 'X'))
        with tempfile.TemporaryDirectory() as tmp:
            update_progressive(first, tmp)
            update_progressive(second, tmp)
            book = load_workbook(history_root(tmp) / 'report.xlsx')
            headers = {cell.value: cell.column for cell in book['URL'][1]}
            self.assertEqual(book['URL'].cell(3, headers['X']).value, 'https://example.com/x/0')
            self.assertEqual(book['URL'].cell(3, headers['Y']).value, 'https://example.com/y/0')
            book.close()

    def test_same_day_rerun_is_retained_but_same_run_id_is_not_duplicated(self):
        with tempfile.TemporaryDirectory() as tmp:
            first = saved_run('first', '2026-10-27T09:00:00+00:00')
            second = saved_run('second', '2026-10-27T12:00:00+00:00', score=51)
            update_progressive(first, tmp)
            update_progressive(second, tmp)
            summary = update_progressive(second, tmp)
            self.assertEqual((summary['days'], summary['runs']), (1, 2))

    def test_portal_download_and_rebuild_use_saved_history(self):
        with tempfile.TemporaryDirectory() as tmp, patch.object(server, 'OUTPUT', Path(tmp)):
            client = server.app.test_client()
            self.assertEqual(client.get('/health').get_json()['progressive_history'], True)
            self.assertEqual(client.get('/progressive/status').get_json()['runs'], 0)
            update_progressive(saved_run('first', '2026-10-27T09:00:00+00:00'), tmp)
            self.assertEqual(client.get('/progressive/status').get_json()['runs'], 1)
            response = client.get('/progressive/report.xlsx')
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.data[:2], b'PK')
            response.close()
            page = client.get('/progressive/report.html')
            self.assertIn(b'Daily Core Web Vitals history', page.data)
            page.close()
            self.assertEqual(client.get('/progressive/invalid').status_code, 404)
            self.assertEqual(client.post('/progressive/rebuild').status_code, 403)
            self.assertEqual(client.post('/progressive/rebuild', headers={'X-Report-Token': server.TOKEN}).get_json()['runs'], 1)


if __name__ == '__main__':
    unittest.main()
