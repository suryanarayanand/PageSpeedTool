"""RSS -> PageSpeed Insights / CrUX -> reproducible Excel and HTML reports."""
import argparse
from datetime import datetime, timedelta, timezone
import html
import json
import math
import os
from pathlib import Path
import random
import re
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
import xml.etree.ElementTree as ET
import sys

if __name__ == "__main__":
    # Keep exception classes shared with the browser collector when launched as a script.
    sys.modules["engine"] = sys.modules[__name__]

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "reports"
CATEGORIES = ["performance", "accessibility", "best-practices", "seo"]
# Report name, PSI field, CrUX field, unit divisor, good maximum, moderate maximum.
FIELD = [
    ("LCP (s)", "LARGEST_CONTENTFUL_PAINT_MS", "largest_contentful_paint", 1000, 2.5, 4),
    ("INP (ms)", "INTERACTION_TO_NEXT_PAINT", "interaction_to_next_paint", 1, 200, 500),
    ("CLS", "CUMULATIVE_LAYOUT_SHIFT_SCORE", "cumulative_layout_shift", 1, .1, .25),
    ("FCP (s)", "FIRST_CONTENTFUL_PAINT_MS", "first_contentful_paint", 1000, 1.8, 3),
    ("TTFB (s)", "EXPERIMENTAL_TIME_TO_FIRST_BYTE", "experimental_time_to_first_byte", 1000, .8, 1.8),
]
LAB = [("Lab FCP (s)", "first-contentful-paint", 1000),
       ("Lab LCP (s)", "largest-contentful-paint", 1000),
       ("Lab TBT (ms)", "total-blocking-time", 1),
       ("Lab CLS", "cumulative-layout-shift", 1),
       ("Lab Speed Index (s)", "speed-index", 1000)]
METRICS = [x[0] for x in FIELD] + ["Performance", "Accessibility", "Best Practices", "SEO"] + [x[0] for x in LAB]
COLORS = {"Good": "DCFCE7", "Needs improvement": "FEF3C7", "Poor": "FEE2E2", "Passed": "DCFCE7", "Failed": "FEE2E2", "Partial": "FEF3C7", "Error": "FEE2E2", "OK": "DCFCE7"}


class FetchError(Exception):
    def __init__(self, message, status=None):
        super().__init__(message)
        self.status = status


def now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def validate_url(url):
    p = urllib.parse.urlsplit(url)
    if p.scheme not in {"http", "https"} or not p.hostname or p.username or p.password:
        raise ValueError("Use an HTTP or HTTPS URL without embedded credentials.")
    return urllib.parse.urlunsplit(p._replace(fragment=""))


def fetch(url, payload=None, timeout=120, attempts=3):
    """Bounded retries; never include request URLs (which may contain keys) in errors."""
    data = json.dumps(payload).encode() if payload is not None else None
    headers = {"User-Agent": "Mozilla/5.0 RSS-PageSpeed-Reporter/1.0", "Accept": "*/*"}
    if data:
        headers["Content-Type"] = "application/json"
    for attempt in range(attempts):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, data=data, headers=headers), timeout=timeout) as response:
                raw = response.read(30_000_001)
                if len(raw) > 30_000_000:
                    raise FetchError("Response exceeds 30 MB limit")
                return raw
        except urllib.error.HTTPError as exc:
            status = exc.code
            messages = {429: "Quota exceeded. Configure a Google API key with available quota and rerun.",
                        403: "Access denied. Check API enablement/key restrictions or feed access.",
                        400: "Request rejected. Check URL and API key.",
                        404: "No data or resource found."}
            error = FetchError(f"HTTP {status}: {messages.get(status, 'Service request failed.')}", status)
            if status not in {429, 500, 502, 503, 504} or attempt == attempts - 1:
                raise error from None
            retry_after = exc.headers.get("Retry-After", "")
            delay = min(float(retry_after), 30) if retry_after.isdigit() else 2 ** (attempt + 1)
        except (urllib.error.URLError, TimeoutError, OSError):
            error = FetchError("Network connection failed or timed out.")
            if attempt == attempts - 1:
                raise error from None
            delay = 2 ** (attempt + 1)
        time.sleep(delay)


def fetch_json(url, payload=None):
    try:
        result = json.loads(fetch(url, payload))
    except (ValueError, UnicodeError):
        raise FetchError("Service returned invalid JSON") from None
    if not isinstance(result, dict):
        raise FetchError("Service returned an unexpected response")
    return result


def parse_feed(raw, feed_url):
    if b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
        raise ValueError("Feeds containing DTD/entity declarations are not supported")
    root = ET.fromstring(raw)
    local = lambda tag: tag.rsplit("}", 1)[-1]
    if local(root.tag) not in {"rss", "RDF", "feed"}:
        raise ValueError("The feed URL did not return RSS or Atom XML")
    found = {}
    for entry in root.iter():
        if local(entry.tag) not in {"item", "entry"}:
            continue
        title, link = "", ""
        for child in entry:
            name = local(child.tag)
            if name == "title":
                title = "".join(child.itertext()).strip()
            if name == "link" and child.get("rel", "alternate") == "alternate" and not link:
                link = child.get("href") or (child.text or "").strip()
        if link:
            try:
                url = validate_url(urllib.parse.urljoin(feed_url, link))
                found.setdefault(url, {"url": url, "title": title})
            except ValueError:
                continue
    return list(found.values())


def number(value, divisor=1):
    try:
        n = float(value) / divisor
        return n if math.isfinite(n) else None
    except (ValueError, TypeError):
        return None


def rating(value, good, moderate):
    if value is None:
        return "Unavailable"
    return "Good" if value <= good else "Needs improvement" if value <= moderate else "Poor"


def field_metrics(metrics, source):
    values = {}
    for name, psi, crux, divisor, _, _ in FIELD:
        item = metrics.get(crux if source == "CrUX" else psi, {})
        val = item.get("percentiles", {}).get("p75") if source == "CrUX" else item.get("percentile")
        # PSI encodes CLS as percentile * 100; CrUX returns the actual decimal.
        values[name] = number(val, 100 if source == "PSI" and name == "CLS" else divisor)
    return values


def assessment(values):
    if any(values.get(name) is None for name, *_ in FIELD[:3]):
        return "Unavailable"
    return "Passed" if all(values[name] <= good for name, _, _, _, good, _ in FIELD[:3]) else "Failed"


def extract_psi(data):
    lh = data.get("lighthouseResult", {})
    error = lh.get("runtimeError") or data.get("error")
    row = {"Status": "Error" if error or not lh else "OK", "Error": "", "Field source": "Unavailable",
           "Field URL": "", "Field period": "", "Final URL": lh.get("finalUrl", ""),
           "Lighthouse version": lh.get("lighthouseVersion", ""), "Measured UTC": lh.get("fetchTime", now()),
           "Warnings": "; ".join(map(str, lh.get("runWarnings", [])))}
    row.update({name: None for name in METRICS})
    if row["Status"] == "Error":
        row["Error"] = str(error.get("message", error.get("code", "Lighthouse failed"))) if isinstance(error, dict) else "Missing Lighthouse result"
    else:
        for name, key in zip(METRICS[5:9], CATEGORIES):
            row[name] = number(lh.get("categories", {}).get(key, {}).get("score"), .01)
        for name, key, divisor in LAB:
            row[name] = number(lh.get("audits", {}).get(key, {}).get("numericValue"), divisor)
        if row["Performance"] is None:
            row["Status"], row["Error"] = "Error", "Lighthouse returned no performance score"
    for key, scope in [("loadingExperience", "URL"), ("originLoadingExperience", "Origin")]:
        experience = data.get(key, {})
        values = field_metrics(experience.get("metrics", {}), "PSI")
        if any(v is not None for v in values.values()):
            row.update(values)
            row["Field source"] = f"PSI / {scope}"
            row["Field URL"] = experience.get("id", "")
            break
    row["CWV assessment"] = assessment(row)
    return row


def add_crux(row, url, strategy, key, raw_dir, stem):
    if not key:
        if row["Field source"] == "Unavailable":
            row["Warnings"] += "; Real-user data unavailable. Configure CRUX_API_KEY to query CrUX directly."
        return
    endpoint = "https://chromeuxreport.googleapis.com/v1/records:queryRecord?" + urllib.parse.urlencode({"key": key})
    parsed = urllib.parse.urlsplit(url)
    origin = f"{parsed.scheme}://{parsed.netloc}"
    for scope, target in [("url", url), ("origin", origin)]:
        try:
            data = fetch_json(endpoint, {scope: target, "formFactor": "PHONE" if strategy == "mobile" else "DESKTOP",
                                        "metrics": [x[2] for x in FIELD]})
            save_json(raw_dir / f"{stem}_crux_{scope}.json", data)
            record = data.get("record", {})
            values = field_metrics(record.get("metrics", {}), "CrUX")
            if not any(v is not None for v in values.values()):
                continue
            row.update(values)
            row["Field source"] = f"CrUX / {scope.title()}"
            row["Field URL"] = record.get("key", {}).get(scope, target)
            row["Field period"] = json.dumps(record.get("collectionPeriod", {}))
            row["CWV assessment"] = assessment(row)
            return
        except FetchError as exc:
            if exc.status == 404:
                continue
            row["Warnings"] += f"; CrUX: {exc}"
            return
    row["Warnings"] += "; CrUX has no real-user data for this URL or origin."


def save_json(path, data):
    temp = path.with_suffix(".tmp")
    temp.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    temp.replace(path)


def validate_config(config):
    size = config.get("sample_size", 1)
    if type(size) is not int or not 1 <= size <= 50:
        raise ValueError("Sample size must be an integer between 1 and 50 per publication")
    feeds = config.get("feeds")
    if not isinstance(feeds, dict) or not feeds or len(feeds) > 20:
        raise ValueError("Provide between 1 and 20 named RSS feeds")
    normalized = {}
    for name, urls in feeds.items():
        if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_-]{0,24}", name) or name.casefold() in {"summary", "results", "url", "urls", "benchmark", "errors", "notes"}:
            raise ValueError("Feed names must be unique short letters/numbers, not reserved report sheet names")
        urls = [urls] if isinstance(urls, str) else urls
        if not isinstance(urls, list) or not 1 <= len(urls) <= 20:
            raise ValueError("Provide between 1 and 20 feed URLs per publication")
        normalized[name] = list(dict.fromkeys(validate_url(url) for url in urls))
    if len({n.casefold() for n in feeds}) != len(feeds):
        raise ValueError("Feed names must be unique ignoring case")
    strategies = config.get("strategies", ["mobile", "desktop"])
    if not isinstance(strategies, list) or not strategies or any(s not in {"mobile", "desktop"} for s in strategies) or len(set(strategies)) != len(strategies):
        raise ValueError("Choose mobile, desktop, or both once")
    collector = config.get("collector", "api")
    if collector not in {"api", "browser"}:
        raise ValueError("Choose website or API collection")
    return {"feeds": normalized, "sample_size": size, "strategies": strategies, "seed": config.get("seed"), "collector": collector}


def run(config, output=OUTPUT, psi_key=None, crux_key=None, progress=None, run_id=None):
    config = validate_config(config)
    run_id = run_id or datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S_") + uuid.uuid4().hex[:8]
    folder = Path(output) / run_id
    raw_dir = folder / "raw"
    raw_dir.mkdir(parents=True, exist_ok=False)
    seed = config.get("seed")
    if seed is None:
        seed = random.SystemRandom().randrange(2 ** 32)
    rng = random.Random(seed)
    state = {"id": run_id, "started": now(), "seed": seed, "config": config,
             "status": "Running", "worker_pid": os.getpid(), "message": "Fetching feeds", "results": [], "samples": [], "errors": []}
    def checkpoint(message):
        state["message"] = message
        save_json(folder / "run.json", state)
        if progress:
            progress(message)
    checkpoint("Fetching RSS feeds")
    for site, feeds in config["feeds"].items():
        pool = {}
        for feed_index, feed in enumerate(feeds):
            checkpoint(f"Fetching {site} feed {feed_index + 1}/{len(feeds)}")
            try:
                raw = fetch(feed, timeout=30)
                items = parse_feed(raw, feed)
                if not items:
                    raise ValueError("Feed contains no usable article links")
                (raw_dir / f"{site}_feed_{feed_index + 1}.xml").write_bytes(raw)
                for item in items:
                    pool.setdefault(item["url"], {"Publication": site, "RSS feed": feed, "URL": item["url"], "Title": item["title"]})
            except (FetchError, ValueError, ET.ParseError) as exc:
                state["errors"].append({"Publication": site, "Stage": "RSS / " + feed, "Error": str(exc)})
        selected = rng.sample(list(pool.values()), min(config["sample_size"], len(pool)))
        if len(selected) < config["sample_size"]:
            state["errors"].append({"Publication": site, "Stage": "Sampling", "Error": f"Requested {config['sample_size']} articles; only {len(selected)} unique links available"})
        state["samples"].extend(selected)
    checkpoint(f"Selected {len(state['samples'])} articles")
    psi_key = psi_key if psi_key is not None else os.getenv("PAGESPEED_API_KEY", "")
    crux_key = crux_key if crux_key is not None else os.getenv("CRUX_API_KEY", "")
    total = len(state["samples"]) * len(config["strategies"])
    for index, sample in enumerate(state["samples"]):
        for strategy in config["strategies"]:
            checkpoint(f"Analyzing {len(state['results']) + 1}/{total}: {sample['Publication']} / {strategy}")
            stem = f"{index + 1:03d}_{sample['Publication']}_{strategy}"
            params = {"url": sample["URL"], "strategy": strategy, "category": CATEGORIES}
            if psi_key:
                params["key"] = psi_key
            try:
                if config["collector"] == "browser":
                    from browser_collector import collect
                    prior = next((r.get('PageSpeed link') for r in state['results'] if r['URL'] == sample['URL'] and r['Status'] == 'OK'), None)
                    row = collect(sample["URL"], strategy, raw_dir, stem, analysis_url=prior)
                else:
                    data = fetch_json("https://www.googleapis.com/pagespeedonline/v5/runPagespeed?" + urllib.parse.urlencode(params, doseq=True))
                    # Do not persist API keys even if echoed by a service.
                    if psi_key:
                        data = json.loads(json.dumps(data).replace(psi_key, "[REDACTED]"))
                    save_json(raw_dir / f"{stem}_psi.json", data)
                    row = extract_psi(data)
            except FetchError as exc:
                row = extract_psi({})
                row["Error"] = str(exc)
            if config["collector"] == "api":
                add_crux(row, sample["URL"], strategy, crux_key, raw_dir, stem)
            row.update(sample)
            row["Device"] = strategy
            row.setdefault("PageSpeed link", "https://pagespeed.web.dev/analysis?" + urllib.parse.urlencode({"url": sample["URL"], "form_factor": strategy}))
            state["results"].append(row)
            checkpoint(f"Collected {len(state['results'])}/{total}")
    ok = sum(r["Status"] == "OK" for r in state["results"])
    has_data = any(any(r.get(k) is not None for k in METRICS) for r in state['results'])
    state["status"] = "Complete" if ok == total and total and not state["errors"] else "Partial" if has_data else "Failed"
    state["finished"] = now()
    write_reports(state, folder)
    checkpoint(f"Report ready: {ok}/{total} successful checks." if state['status'] == 'Complete' else
               f"Partial report: {ok}/{total} complete checks. Available measurements are saved; review the missing values." if has_data else
               "No measurements collected. See failure details and retry.")
    return folder


def display(value):
    if value is None:
        return "N/A"
    if isinstance(value, float):
        return f"{value:.3f}".rstrip("0").rstrip(".")
    return str(value)


def safe_cell(value):
    if isinstance(value, str):
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)[:32000]
        if value.lstrip().startswith(("=", "+", "-", "@")):
            value = "'" + value
    return "N/A" if value is None else value


def table_sheet(wb, name, headers, rows):
    sheet = wb.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append([safe_cell(v) for v in row])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[1].height = 36
    for i, heading in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(i)].width = 55 if "URL" in heading or "Error" in heading or "Title" in heading else 22
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, float):
                cell.number_format = "0.###"
            if isinstance(cell.value, str) and cell.value in COLORS:
                cell.fill = PatternFill("solid", fgColor=COLORS[cell.value])
    return sheet


def summary_rows(state):
    result = []
    for site in state["config"]["feeds"]:
        for device in state["config"]["strategies"]:
            rows = [r for r in state["results"] if r["Publication"] == site and r["Device"] == device]
            scores = [r["Performance"] for r in rows if r["Performance"] is not None]
            result.append([site, device, len(rows), sum(r["Status"] == "OK" for r in rows),
                           sum(scores) / len(scores) if scores else None])
    return result


def write_reports(state, folder):
    from excel_layout import reference_workbook
    wb = reference_workbook(state, METRICS, FIELD, safe_cell)
    summary_headers = ["Publication", "Device", "Audits", "Complete checks", "Average performance"]
    summaries = summary_rows(state)
    table_sheet(wb, "Summary", summary_headers, summaries)
    headers = ["Publication", "URL", "Title", "Device", "Status", "Field source", "Field URL", "Field period"] + METRICS + ["Measured UTC", "Final URL", "Lighthouse version", "PageSpeed link", "Warnings", "Error", "Collection method", "PageSpeed captured time"]
    clean_warning = lambda r: r.get('Warnings', '').replace(' CWV assessment is taken from the website.', '')
    detailed = table_sheet(wb, "Results", headers, [[clean_warning(r) if k == 'Warnings' else r.get(k) for k in headers] for r in state["results"]])
    for cells, result in zip(detailed.iter_rows(min_row=2), state["results"]):
        for metric, _, _, _, good, moderate in FIELD:
            cell = cells[headers.index(metric)]
            color = COLORS.get(rating(result.get(metric), good, moderate))
            if color:
                cell.fill = PatternFill("solid", fgColor=color)
        for metric in METRICS[5:9]:
            value = result.get(metric)
            if value is not None:
                cells[headers.index(metric)].fill = PatternFill("solid", fgColor=COLORS["Good" if value >= 90 else "Needs improvement" if value >= 50 else "Poor"])
    run_date = datetime.fromisoformat(state['started']).astimezone(timezone(timedelta(hours=5, minutes=30))).date()
    url_details = table_sheet(wb, "URLs", ["Date", "Publication", "RSS feed", "URL", "Title"], [[run_date] + [s[k] for k in ["Publication", "RSS feed", "URL", "Title"]] for s in state["samples"]])
    for row in url_details.iter_rows(min_row=2, max_col=1):
        row[0].number_format = 'd-mmm-yy'
    errors = state["errors"] + [{"Publication": r["Publication"], "Stage": r["Device"] + " / " + r["URL"], "Error": r["Error"]} for r in state["results"] if r["Status"] != "OK"]
    table_sheet(wb, "Errors", ["Publication", "Stage", "Error"], [[e[k] for k in ["Publication", "Stage", "Error"]] for e in errors])
    notes = [[name, good, moderate, "Above moderate maximum"] for name, _, _, _, good, moderate in FIELD]
    notes += [["Collection", state['config'].get('collector', 'api'), "Website values use displayed precision", "Numeric metrics and scores"], ["Scores", "90–100", "50–89", "0–49"], ["Run status", state["status"], "UTC started", state["started"]],
              ["Random seed", state["seed"], "Run ID", state["id"]],
              ["Field data", "75th percentile, rolling 28-day real-user data", "Origin data describes the whole origin", "N/A means unavailable, never zero"],
              ["Lighthouse", "Single lab run; results may vary", "TBT is not INP", "Field and lab values are separate"],
              ["Sampling", "Without replacement within each feed", "Different runs can repeat articles", "Raw feeds and API responses saved beside reports"]]
    notes.append(['Dates', 'Publication and URL sheets use India time', 'One row per sampled article', 'Publication row 4 corresponds to URL row 2; subsequent rows follow the same order'])
    table_sheet(wb, "Notes", ["Metric / note", "Good maximum / value", "Moderate maximum / detail", "Poor / note"], notes)
    wb.save(folder / "report.xlsx")
    esc = lambda v: html.escape(display(v), quote=True)
    def html_table(heads, rows):
        return '<div class="scroll"><table><thead><tr>' + ''.join(f'<th>{esc(h)}</th>' for h in heads) + '</tr></thead><tbody>' + ''.join('<tr>' + ''.join(f'<td>{esc(v)}</td>' for v in row) + '</tr>' for row in rows) + '</tbody></table></div>'
    cards = []
    for r in state["results"]:
        cells = []
        for metric in METRICS:
            grade = ""
            for name, _, _, _, good, moderate in FIELD:
                if metric == name:
                    grade = rating(r.get(metric), good, moderate)
            cells.append([metric, r.get(metric), grade])
        cards.append(f'<article class="audit"><h3>{esc(r["Publication"])} · {esc(r["Device"])} · {esc(r["Status"])}</h3><a href="{esc(r["URL"])}">{esc(r["Title"] or r["URL"])}</a><p>Field data: {esc(r["Field source"])}</p><p>Field URL: {esc(r["Field URL"])}<br>Period: {esc(r["Field period"])}</p>' + html_table(["Metric", "Value", "Field rating"], cells) + f'<p>{esc(clean_warning(r))}</p><p class="error">{esc(r["Error"])}</p><a href="{esc(r["PageSpeed link"])}">Open in PageSpeed Insights</a></article>')
    page = '<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>RSS PageSpeed report</title><style>' + STYLE + '</style><main>'
    if not any(any(r.get(k) is not None for k in METRICS) for r in state['results']):
        page += '<div class="panel error"><h1>No measurements collected</h1><p>This run failed. The tables below contain diagnostic details, not valid performance measurements. Retry the analysis after addressing the listed errors.</p></div>'
    page += f'<h1>RSS PageSpeed report</h1><p>{esc(state["status"])} · {esc(state["started"])} · Seed {esc(state["seed"])}</p><a class="button" href="report.xlsx">Download Excel</a><p>Real-user metrics are p75 over a rolling 28-day window. Origin data represents the whole site. N/A means unavailable. Lighthouse metrics are individual lab runs.</p><h2>Publication summary</h2>'
    page += html_table(summary_headers, summaries)
    page += '<h2>Article results</h2><label>Filter by publication, device or URL <input id="filter" type="search" placeholder="e.g. mobile"></label><div class="grid">' + ''.join(cards) + '</div>'
    page += '<h2>Feed and audit errors</h2>' + html_table(["Publication", "Stage", "Error"], [[e[k] for k in ["Publication", "Stage", "Error"]] for e in errors])
    page += '<p>Source: <a href="https://developers.google.com/speed/docs/insights/v5/about">PageSpeed Insights</a> · <a href="https://developer.chrome.com/docs/crux/api">CrUX API</a></p></main><script>document.getElementById("filter").addEventListener("input",function(){const q=this.value.toLowerCase();document.querySelectorAll(".audit").forEach(a=>a.hidden=!a.textContent.toLowerCase().includes(q));});</script></html>'
    (folder / "report.html").write_text(page, encoding="utf-8")
    save_json(folder / 'report_layout.json', {'version': 2, 'excel': 'reference', 'cwv_assessment': False})


STYLE = """*{box-sizing:border-box}body{margin:0;background:#f1f5f9;color:#17324d;font:15px system-ui,sans-serif}main{max-width:1450px;margin:auto;padding:32px}h1{font-size:34px}h2{margin-top:32px}a{color:#0369a1;overflow-wrap:anywhere}p{line-height:1.6;overflow-wrap:anywhere}.scroll{overflow:auto}table{border-collapse:collapse;width:100%;background:white}td,th{padding:10px;text-align:left;border-bottom:1px solid #e2e8f0;min-width:100px}th{background:#17324d;color:white}.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,440px),1fr));gap:20px;margin-top:20px}article,.panel{background:white;border:1px solid #cbd5e1;border-radius:12px;padding:22px}input,textarea,select{font:inherit;padding:10px;border:1px solid #94a3b8;border-radius:6px;max-width:100%}label{display:block;margin:14px 0}textarea{width:100%;height:165px}button,.button{display:inline-block;border:0;border-radius:6px;background:#0369a1;color:white;padding:12px 18px;text-decoration:none;font:inherit;cursor:pointer}button:disabled{opacity:.5}.error{color:#b91c1c}[hidden]{display:none!important}code{overflow-wrap:anywhere}"""


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=ROOT / "config.json")
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--sample-size", type=int)
    parser.add_argument("--seed", type=int)
    args = parser.parse_args()
    config = json.loads(args.config.read_text(encoding="utf-8-sig"))
    if args.sample_size is not None:
        config["sample_size"] = args.sample_size
    if args.seed is not None:
        config["seed"] = args.seed
    path = run(config, args.output, progress=lambda message: print(message, flush=True))
    print(path)
    raise SystemExit(0 if json.loads((path / "run.json").read_text(encoding="utf-8"))["status"] == "Complete" else 2)
