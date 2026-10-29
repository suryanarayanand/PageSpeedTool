from io import BytesIO
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

import engine
from browser_collector import LIGHTHOUSE_READY, collect, parse_visible
from progressive_report import history_root


PARTIAL_PAGE = '''Core Web Vitals Assessment:
Passed
Largest Contentful Paint (LCP)
1.6 s
Interaction to Next Paint (INP)
110 ms
Cumulative Layout Shift (CLS)
0.03
First Contentful Paint (FCP)
1 s
Time to First Byte (TTFB)
0.5 s
Diagnose performance issues
Performance
74
Accessibility
54
Best Practices
92
SEO
1/2
Agentic Browsing
First Contentful Paint
Error
Largest Contentful Paint
Error
Total Blocking Time
Error
Cumulative Layout Shift
Error
Speed Index
Error
'''
SCORES = {'Performance': None, 'Accessibility': 74, 'Best Practices': 54, 'SEO': 92}


def partial_row():
    return parse_visible(PARTIAL_PAGE, 'https://example.com/article', 'desktop', 'Origin',
                         'https://pagespeed.web.dev/analysis/example/report',
                         category_scores=SCORES, terminal_errors=['Performance'])


class PartialLighthouseTests(unittest.TestCase):
    def test_performance_error_retains_field_data_and_other_category_scores(self):
        row = partial_row()
        self.assertEqual(row['Status'], 'Partial')
        self.assertEqual([row[key] for key in ('LCP (s)', 'INP (ms)', 'CLS', 'FCP (s)', 'TTFB (s)')],
                         [1.6, 110, .03, 1, .5])
        self.assertEqual(row['Field source'], 'PageSpeed website / Origin')
        self.assertEqual(row['Field URL'], 'https://example.com')
        self.assertIsNone(row['Performance'])
        self.assertEqual([row[key] for key in ('Accessibility', 'Best Practices', 'SEO')], [74, 54, 92])
        self.assertTrue(row['Lighthouse terminal error'])
        with patch('browser_collector.collect_once', return_value=row) as audit:
            self.assertEqual(collect('https://example.com/article', 'desktop', Path('.'), 'audit'), row)
            audit.assert_called_once()

    def test_partial_data_reaches_individual_and_progressive_excel(self):
        rss = b'<rss><channel><item><title>Article</title><link>https://example.com/article</link></item></channel></rss>'
        config = {'feeds': {'BL': 'https://example.com/feed'}, 'sample_size': 1,
                  'strategies': ['desktop'], 'collector': 'browser'}
        with tempfile.TemporaryDirectory() as tmp, patch('engine.fetch', return_value=rss), \
                patch('browser_collector.collect', return_value=partial_row()):
            folder = engine.run(config, tmp)
            state = json.loads((folder / 'run.json').read_text())
            self.assertEqual(state['status'], 'Partial')
            for path in (folder / 'report.xlsx', history_root(tmp) / 'report.xlsx'):
                book = load_workbook(path)
                self.assertEqual([book['BL'].cell(4, col).value for col in range(16, 21)], [1.6, 110, .03, 1, .5])
                self.assertEqual(book['BL']['U4'].value, 'N/A')
                self.assertEqual([book['BL'].cell(4, col).value for col in range(22, 25)], [74, 54, 92])
                book.close()
            self.assertEqual(json.loads((history_root(tmp) / 'summary.json').read_text())['measured_checks'], 1)

    @unittest.skipUnless(shutil.which('node'), 'Node is needed to check the browser readiness predicate')
    def test_browser_wait_distinguishes_finished_error_loading_and_other_device(self):
        fixture = r'''
const fs = require('fs');
const input = JSON.parse(fs.readFileSync(0, 'utf8'));
const ready = eval('(' + input.source + ')');
function evaluate(state) {
  const gauges = Object.entries(state.scores).map(([label, score]) => ({
    getClientRects: () => state.hidden === label ? [] : [1],
    querySelector: selector => selector === '.lh-gauge__label' ? {textContent: label} : {
      textContent: score === null || score === 'loading' ? '' : String(score),
      classList: {contains: () => score === null}
    }
  }));
  const metrics = Array.from({length: state.metrics}, () => ({getClientRects: () => [1], innerText: '1.2 s'}));
  const panel = {innerText: '', querySelectorAll: selector =>
    selector === '.lh-gauge__wrapper' ? gauges : selector === '.lh-metric__value' ? metrics : []};
  global.document = {getElementById: () => ({getAttribute: () => state.selected ? 'true' : 'false'}),
                     querySelector: () => panel};
  return ready('desktop');
}
process.stdout.write(JSON.stringify(input.cases.map(evaluate)));
'''
        cases = [
            {'scores': SCORES, 'metrics': 0, 'selected': True},
            {'scores': dict(SCORES, Performance='loading'), 'metrics': 0, 'selected': True},
            {'scores': SCORES, 'metrics': 0, 'selected': False},
            {'scores': dict(SCORES, Performance=80), 'metrics': 5, 'selected': True},
            {'scores': SCORES, 'metrics': 0, 'selected': True, 'hidden': 'Performance'},
        ]
        result = subprocess.run([shutil.which('node'), '-e', fixture],
                                input=json.dumps({'source': LIGHTHOUSE_READY, 'cases': cases}),
                                text=True, capture_output=True, check=True)
        values = json.loads(result.stdout)
        self.assertEqual(values[0]['terminal_errors'], ['Performance'])
        self.assertEqual(values[0]['scores']['Accessibility'], 74)
        self.assertFalse(values[1])
        self.assertFalse(values[2])
        self.assertEqual(values[3]['scores']['Performance'], 80)
        self.assertEqual(values[3]['terminal_errors'], [])
        self.assertFalse(values[4])


if __name__ == '__main__':
    unittest.main()
