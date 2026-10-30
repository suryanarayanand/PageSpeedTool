from html.parser import HTMLParser
import json
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import unittest

from openpyxl import load_workbook

from cloud_audit import audit_config, summarize_audit
from pages_history import push_history, restore_history
from progressive_report import history_root, update_progressive
from publish_static import publish_site
from test_progressive import saved_run


class Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []

    def handle_starttag(self, tag, attributes):
        if tag == 'a':
            self.links.extend(value for name, value in attributes if name == 'href')


class StaticCloudTests(unittest.TestCase):
    def test_workflow_summary_distinguishes_failed_measurements_and_preserves_field_data(self):
        from contextlib import redirect_stdout
        from io import StringIO
        from unittest.mock import patch
        state = saved_run('limited', '2026-10-30T06:30:00+00:00', failed=True)
        for row in state['results']:
            row['Error'] = 'PageSpeed returned HTTP 429. <service>'
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / 'summary.md'
            captured = StringIO()
            with patch.dict('os.environ', {'GITHUB_STEP_SUMMARY': str(target)}), redirect_stdout(captured):
                self.assertEqual(summarize_audit(state), 0)
            text = target.read_text(encoding='utf-8')
            self.assertIn('Measurement status: Failed', text)
            self.assertIn('0/2 checks', text)
            self.assertIn('HTTP 429. &lt;service&gt;', text)
            self.assertIn('No new measurements were collected', captured.getvalue())
            state['status'] = 'Partial'
            state['results'][0]['LCP (s)'] = 2.1
            with patch.dict('os.environ', {'GITHUB_STEP_SUMMARY': str(target)}), redirect_stdout(StringIO()):
                self.assertEqual(summarize_audit(state), 1)
            self.assertIn('1/2 checks', target.read_text(encoding='utf-8'))

    def test_cloud_audit_is_article_count_based_and_can_use_browser_without_api_key(self):
        config = audit_config({'ARTICLE_COUNT': '2', 'PUBLICATIONS': 'BL', 'DEVICES': 'mobile', 'COLLECTOR': 'auto'})
        self.assertEqual(config['sample_size'], 2)
        self.assertEqual(list(config['feeds']), ['BL'])
        self.assertEqual(config['strategies'], ['mobile'])
        self.assertEqual(config['collector'], 'browser')
        self.assertEqual(audit_config({'PAGESPEED_API_KEY': 'server-key'})['collector'], 'api')
        with self.assertRaises(ValueError):
            audit_config({'ARTICLE_COUNT': '0'})
        with self.assertRaises(ValueError):
            audit_config({'PUBLICATIONS': 'FL'})

    def test_static_portal_supports_project_urls_daily_filter_and_excel_downloads(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / 'reports'
            site = Path(tmp) / 'site'
            update_progressive(saved_run('first', '2026-10-27T09:00:00+00:00', score=90), output)
            update_progressive(saved_run('second', '2026-10-28T09:00:00+00:00', score=42), output)
            result = publish_site(site, output, 'owner/PageSpeedTool')
            self.assertEqual(result['runs'], 2)
            page = (site / 'index.html').read_text(encoding='utf-8')
            self.assertIn('data/summary.json', page)
            self.assertIn('data/runs.json', page)
            self.assertIn('date-filter', page)
            self.assertNotIn('127.0.0.1', page)
            self.assertNotIn("fetch('/", page)
            self.assertIn('https://github.com/owner/PageSpeedTool/actions/workflows/daily-pages.yml', page)
            links = Links()
            links.feed(page)
            for link in links.links:
                if not link.startswith(('https://', '#')):
                    self.assertFalse(link.startswith('/'), link)
                    self.assertTrue((site / link).is_file(), link)
            runs = json.loads((site / 'data/runs.json').read_text(encoding='utf-8'))
            self.assertEqual([row['report_date'] for row in runs], ['2026-10-28', '2026-10-27'])
            self.assertEqual([row['measured_checks'] for row in runs], [2, 2])
            self.assertNotIn('worker_pid', json.dumps(runs))
            self.assertTrue((site / 'reports/first/report.xlsx').is_file())
            self.assertIn('href="../index.html"', (site / 'progressive/report.html').read_text(encoding='utf-8'))
            book = load_workbook(site / 'progressive/report.xlsx')
            self.assertEqual([book['TH'].cell(row, 7).value for row in (4, 5)], [90, 42])
            book.close()
            publish_site(site, output, 'owner/PageSpeedTool')
            self.assertEqual(json.loads((site / 'data/summary.json').read_text(encoding='utf-8'))['runs'], 2)
            if shutil.which('node'):
                script = re.search(r'<script>(.*?)</script>', page, re.S).group(1)
                subprocess.run([shutil.which('node'), '--check'], input=script, text=True, capture_output=True, check=True)

    def test_static_export_does_not_replace_unrelated_user_files(self):
        with tempfile.TemporaryDirectory() as tmp:
            destination = Path(tmp) / 'existing'
            destination.mkdir()
            (destination / 'user.txt').write_text('keep this', encoding='utf-8')
            with self.assertRaises(ValueError):
                publish_site(destination, Path(tmp) / 'reports')
            self.assertEqual((destination / 'user.txt').read_text(encoding='utf-8'), 'keep this')

    @unittest.skipUnless(shutil.which('node'), 'Node is needed to check portal refresh behavior')
    def test_portal_fetches_fresh_history_versions_downloads_and_refreshes_open_report(self):
        with tempfile.TemporaryDirectory() as tmp:
            output, site = Path(tmp) / 'reports', Path(tmp) / 'site'
            state = saved_run('failed', '2026-10-30T06:30:00+00:00', failed=True)
            for row in state['results']:
                row['Error'] = 'PageSpeed service returned HTTP 429.'
            update_progressive(state, output)
            publish_site(site, output, 'owner/PageSpeedTool')
            page = (site / 'index.html').read_text(encoding='utf-8')
            progressive = (site / 'progressive/report.html').read_text(encoding='utf-8')
            summary = json.loads((site / 'data/summary.json').read_text(encoding='utf-8'))
            runs = json.loads((site / 'data/runs.json').read_text(encoding='utf-8'))
            self.assertEqual(runs[0]['first_error'], 'PageSpeed service returned HTTP 429.')
            source = re.search(r'<script>(.*?)</script>', page, re.S).group(1)
            progressive_source = re.search(r'<script>(.*?)</script>', progressive, re.S).group(1)
            fixture = r'''
const fs=require('fs'),input=JSON.parse(fs.readFileSync(0,'utf8'));
function element(){return {children:[],value:'',getAttribute(key){return this[key]??null;},
  append(...items){this.children.push(...items);},replaceChildren(...items){this.children=[...items];},addEventListener(){}};}
const elements=Object.fromEntries(['date-filter','runs','summary','downloads','progressive-html','progressive-excel'].map(id=>[id,element()]));
elements['progressive-html'].href='progressive/report.html';elements['progressive-excel'].href='progressive/report.xlsx';
global.document={getElementById:id=>elements[id],createElement:()=>element()};
global.window={location:{href:'https://owner.github.io/PageSpeedTool/'}};
let intervalCallback,delay;const requests=[];
global.setInterval=(callback,ms)=>{intervalCallback=callback;delay=ms;};
global.fetch=async(url,options)=>{requests.push({url,cache:options.cache});return {ok:true,json:async()=>url.includes('summary.json')?input.summary:input.runs};};
function textNodes(node){return [node.textContent,...node.children.flatMap(textNodes)].filter(value=>typeof value==='string');}
async function main(){
  eval(input.source);await new Promise(resolve=>setImmediate(resolve));
  const dashboard={requests:[...requests],delay,html:elements['progressive-html'].href,excel:elements['progressive-excel'].href,
    error:textNodes(elements.runs).find(text=>text.includes('HTTP 429')),
    reportLinks:elements.runs.children[0].children.at(-1).children.map(node=>node.href)};
  let redirected;
  window.location={href:'https://owner.github.io/PageSpeedTool/progressive/report.html',replace:url=>{redirected=url;}};
  global.fetch=async(url,options)=>({ok:true,json:async()=>({...input.summary,updated:'new-publication'})});
  eval(input.progressiveSource);await intervalCallback();
  process.stdout.write(JSON.stringify({dashboard,redirected}));
}
main().catch(error=>{process.stderr.write(error.stack);process.exitCode=1;});
'''
            result = subprocess.run([shutil.which('node'), '-e', fixture],
                                    input=json.dumps({'source': source, 'progressiveSource': progressive_source,
                                                      'summary': summary, 'runs': runs}),
                                    text=True, encoding='utf-8', capture_output=True, check=True)
            observed = json.loads(result.stdout)
            from urllib.parse import parse_qs, urlsplit
            for request in observed['dashboard']['requests']:
                url = urlsplit(request['url'])
                self.assertTrue(url.path.startswith('/PageSpeedTool/data/'))
                self.assertTrue(parse_qs(url.query)['_refresh'][0])
                self.assertEqual(request['cache'], 'no-store')
            for target in [observed['dashboard']['html'], observed['dashboard']['excel']] + observed['dashboard']['reportLinks']:
                self.assertEqual(parse_qs(urlsplit(target).query)['v'], [summary['updated']])
            self.assertEqual(observed['dashboard']['error'], 'PageSpeed service returned HTTP 429.')
            self.assertEqual(observed['dashboard']['delay'], 60000)
            self.assertEqual(urlsplit(observed['redirected']).path, '/PageSpeedTool/progressive/report.html')
            self.assertEqual(parse_qs(urlsplit(observed['redirected']).query)['v'], ['new-publication'])

    def test_history_branch_retains_days_without_changing_the_source_branch(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            remote, repository = root / 'origin.git', root / 'source'
            subprocess.run(['git', 'init', '--bare', str(remote)], capture_output=True, check=True)
            subprocess.run(['git', 'init', str(repository)], capture_output=True, check=True)
            def git(*args):
                return subprocess.run(['git', '-C', str(repository), *args], capture_output=True, text=True, check=True).stdout.strip()
            (repository / 'source.txt').write_text('application source', encoding='utf-8')
            git('add', 'source.txt')
            git('-c', 'user.name=Test', '-c', 'user.email=test@example.com', '-c', 'commit.gpgsign=false', 'commit', '-m', 'Source')
            git('remote', 'add', 'origin', str(remote))
            source_head = git('rev-parse', 'HEAD')
            output = root / 'first-runner/reports'
            self.assertEqual(restore_history(repository, output, seed=root / 'empty-seed'), 0)
            state = saved_run('first', '2026-10-27T09:00:00+00:00')
            state['worker_pid'] = 1234
            state['config']['psi_key'] = 'private-secret'
            update_progressive(state, output)
            self.assertTrue(push_history(repository, output)['changed'])
            self.assertFalse(push_history(repository, output)['changed'])
            second_output = root / 'second-runner/reports'
            self.assertEqual(restore_history(repository, second_output), 1)
            restored = (history_root(second_output) / 'runs/first.json').read_text(encoding='utf-8')
            self.assertNotIn('worker_pid', restored)
            self.assertNotIn('private-secret', restored)
            update_progressive(saved_run('second', '2026-10-28T09:00:00+00:00'), second_output)
            self.assertEqual(push_history(repository, second_output)['runs'], 2)
            self.assertEqual(restore_history(repository, output), 2)
            self.assertEqual(git('rev-parse', 'HEAD'), source_head)
            self.assertEqual((repository / 'source.txt').read_text(encoding='utf-8'), 'application source')
            self.assertEqual(git('status', '--porcelain'), '')


if __name__ == '__main__':
    unittest.main()
