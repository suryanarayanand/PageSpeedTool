"""Populate the reference workbook layout with current, traceable measurements."""
from copy import copy
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TEMPLATE = Path(__file__).parent / 'templates' / 'reference_report.xlsx'


def reference_workbook(state, metrics, field_definitions, safe_cell):
    book = load_workbook(TEMPLATE)
    date = datetime.fromisoformat(state['started']).astimezone(timezone(timedelta(hours=5, minutes=30))).date()
    publications = ['TH', 'BL', 'FL', 'SS']
    for site in state['config']['feeds']:
        if site not in publications:
            sheet = book.copy_worksheet(book['TH'])
            sheet.title = site
            publications.append(site)
    samples = {site: [s for s in state['samples'] if s['Publication'] == site] for site in publications}
    audits = {(r['Publication'], r['URL'], r['Device']): r for r in state['results']}
    urls = book['URL']
    for col, site in enumerate(publications, 2):
        urls.cell(1, col, site)
        urls.column_dimensions[get_column_letter(col)].width = 48
    urls.column_dimensions['A'].width = 14
    urls.freeze_panes = 'B2'
    for index in range(max((len(v) for v in samples.values()), default=0)):
        row = index + 2
        urls.cell(row, 1, date).number_format = 'd-mmm-yy'
        urls.row_dimensions[row].height = 48
        for col, site in enumerate(publications, 2):
            if index < len(samples[site]):
                sample = samples[site][index]
                cell = urls.cell(row, col, sample['URL'])
                cell.hyperlink = sample['URL']
                cell.font = Font(name='Aptos Narrow', size=11, color='0563C1', underline='single')
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                cell.comment = Comment(sample['Title'] + '\nRSS feed: ' + sample['RSS feed'], 'PageSpeed Reporter')
    urls.auto_filter.ref = f'A1:{get_column_letter(len(publications) + 1)}{max(2, urls.max_row)}'
    for site in publications:
        sheet = book[site]
        # Date + 14 mobile metrics + 14 desktop metrics exactly match the old file.
        styles = [copy(sheet.cell(4, col)._style) for col in range(1, 30)]
        sheet.freeze_panes = 'B4'
        sheet.sheet_view.zoomScale = 80
        sheet.row_dimensions[3].height = 34
        sheet.column_dimensions['A'].width = 13
        sheet.print_title_rows = '1:3'
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        if not samples[site]:
            sheet.merge_cells('A5:AC5')
            sheet['A5'] = 'Not selected / no RSS feed configured for this run.' if site not in state['config']['feeds'] else 'No articles could be collected. See Errors for details.'
            sheet['A5'].font = Font(name='Aptos Narrow', size=11, italic=True, color='666666')
        for index, sample in enumerate(samples[site]):
            row = index + 4
            for col in range(1, 30):
                cell = sheet.cell(row, col)
                cell._style = copy(styles[col - 1])
                cell.font = Font(name='Aptos Narrow', size=11, color='000000')
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.fill = PatternFill(fill_type=None)
            sheet.row_dimensions[row].height = 22
            date_cell = sheet.cell(row, 1, date)
            date_cell.number_format = 'd-mmm-yy'
            url_col = get_column_letter(publications.index(site) + 2)
            date_cell.hyperlink = f"#'URL'!{url_col}{index + 2}"
            date_cell.comment = Comment(f"Article {index + 1}: {sample['Title']}\n{sample['URL']}\nSee URL!{url_col}{index + 2}. Dates use India time.", 'PageSpeed Reporter')
            for device, start in [('mobile', 2), ('desktop', 16)]:
                audit = audits.get((site, sample['URL'], device), {})
                for offset, name in enumerate(metrics):
                    value = audit.get(name)
                    cell = sheet.cell(row, start + offset, safe_cell(value))
                    cell.number_format = '0' if name in {'INP (ms)', 'Lab TBT (ms)', 'Performance', 'Accessibility', 'Best Practices', 'SEO'} else '0.###'
                    missing_status = 'Not completed (run stopped)' if state.get('status') == 'Stopped' else 'Not requested'
                    detail = f"{sample['URL']}\n{device.title()} | {audit.get('Status', missing_status)}\nField source: {audit.get('Field source', 'Not collected')}\nField URL: {audit.get('Field URL', '')}\n{audit.get('Error', '')}"
                    cell.comment = Comment(detail, 'PageSpeed Reporter')
                    if value is None:
                        cell.font = Font(name='Aptos Narrow', size=11, italic=True, color='777777')
                        cell.fill = PatternFill('solid', fgColor='F2F2F2')
                    elif offset < 5:
                        good, moderate = field_definitions[offset][4:6]
                        color = 'C6EFCE' if value <= good else 'FFEB9C' if value <= moderate else 'FFC7CE'
                        cell.fill = PatternFill('solid', fgColor=color)
                    elif offset < 9:
                        cell.fill = PatternFill('solid', fgColor='C6EFCE' if value >= 90 else 'FFEB9C' if value >= 50 else 'FFC7CE')
        sheet.print_area = f'A1:AC{max(4, sheet.max_row)}'
    # Keep the original primary tabs first. Diagnostic tables follow them.
    book._sheets = [book['BenchMark']] + [book[site] for site in publications] + [book['URL']]
    book.active = book.sheetnames.index(next((site for site in publications if samples[site]), 'TH'))
    return book
