"""Extract only the reference workbook's presentation, without historical results."""
from copy import copy
from pathlib import Path
from openpyxl import Workbook, load_workbook


def create(source, target):
    reference = load_workbook(source)
    book = Workbook()
    book.remove(book.active)
    book.loaded_theme = reference.loaded_theme
    for name in ['BenchMark', 'TH', 'BL', 'FL', 'SS', 'URL']:
        old = reference[name]
        sheet = book.create_sheet(name)
        rows, columns = (13, 10) if name == 'BenchMark' else (2, 5) if name == 'URL' else (4, 29)
        for row in old.iter_rows(min_row=1, max_row=rows, max_col=columns):
            for original in row:
                cell = sheet.cell(original.row, original.column)
                if name == 'BenchMark' or original.row <= (1 if name == 'URL' else 3):
                    cell.value = original.value
                cell.font = copy(original.font)
                cell.fill = copy(original.fill)
                cell.border = copy(original.border)
                cell.alignment = copy(original.alignment)
                cell.number_format = original.number_format
                cell.protection = copy(original.protection)
        for merged in old.merged_cells.ranges:
            if merged.max_row <= rows and merged.max_col <= columns:
                sheet.merge_cells(str(merged))
        for key, dimension in old.column_dimensions.items():
            if dimension.min <= columns:
                sheet.column_dimensions[key] = copy(dimension)
        for index in range(1, rows + 1):
            sheet.row_dimensions[index] = copy(old.row_dimensions[index])
        sheet.sheet_format = copy(old.sheet_format)
        sheet.sheet_properties = copy(old.sheet_properties)
    target.parent.mkdir(parents=True, exist_ok=True)
    book.save(target)


if __name__ == '__main__':
    create(Path(r'C:\Users\12503\Downloads\TH,BL,FL,SS- Core Web Vitals.xlsx'),
           Path(__file__).parent / 'templates' / 'reference_report.xlsx')
