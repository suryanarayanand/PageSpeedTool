import json
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import MagicMock, patch
from openpyxl import load_workbook
import engine
import server
from browser_collector import collect, collect_once, parse_visible, visible_number

RSS = b'<rss><channel><item><title>=bad &lt;script&gt;</title><link>https://example.com/a#top</link></item><item><link>https://example.com/a</link></item><item><link>https://example.com/b</link></item></channel></rss>'


def response():
    return {"loadingExperience": {"id": "https://example.com/a", "metrics": {
        "LARGEST_CONTENTFUL_PAINT_MS": {"percentile": 2500},
        "INTERACTION_TO_NEXT_PAINT": {"percentile": 200},
        "CUMULATIVE_LAYOUT_SHIFT_SCORE": {"percentile": 10}}},
        "lighthouseResult": {"finalUrl": "https://example.com/a", "categories": {
            k: {"score": .9} for k in engine.CATEGORIES}, "audits": {
                "first-contentful-paint": {"numericValue": 1250},
                "total-blocking-time": {"numericValue": 500}}}}


class ReporterTests(unittest.TestCase):
    def test_mac_uses_headless_chrome_and_reuses_a_dedicated_profile(self):
        import browser_collector
        from playwright.sync_api import Error as BrowserError
        with tempfile.TemporaryDirectory() as tmp, patch.object(browser_collector, 'ROOT', Path(tmp)), \
                patch('browser_collector.sys.platform', 'darwin'), patch('playwright.sync_api.sync_playwright') as start:
            launch = start.return_value.__enter__.return_value.chromium.launch_persistent_context
            launch.side_effect = BrowserError('Launch did not complete')
            for _ in range(2):
                with self.assertRaises(engine.FetchError):
                    collect_once('https://example.com/a', 'mobile', Path(tmp), 'audit')
            profiles = [call.kwargs['user_data_dir'] for call in launch.call_args_list]
            self.assertEqual(profiles, [str(Path(tmp) / 'browser_profile')] * 2)
            self.assertTrue(Path(profiles[0]).is_dir())
            self.assertTrue(launch.call_args.kwargs['headless'])
            self.assertTrue(launch.call_args.kwargs['chromium_sandbox'])

    def test_blocked_pagespeed_request_is_identified_and_network_keys_are_not_saved(self):
        from types import SimpleNamespace
        context, page, tab = MagicMock(), MagicMock(), MagicMock()
        context.pages = [page]
        tab.is_visible.return_value = False
        page.locator.side_effect = lambda selector: tab if selector == '#mobile_tab' else page
        page.inner_text.return_value = 'Unable to resolve https://example.com/a. Try checking the URL for validity.'
        def event(name, callback):
            if name == 'requestfailed':
                callback(SimpleNamespace(url='https://content-pagespeedonline.googleapis.com/v5/test?key=private-key',
                                         failure='net::ERR_BLOCKED_BY_CLIENT'))
        page.on.side_effect = event
        with tempfile.TemporaryDirectory() as tmp, patch('playwright.sync_api.sync_playwright') as start, patch('browser_collector.sys.platform', 'linux'):
            launch = start.return_value.__enter__.return_value.chromium.launch_persistent_context
            launch.return_value = context
            with self.assertRaisesRegex(engine.FetchError, 'managed browser extension blocked') as raised:
                collect('https://example.com/a', 'mobile', Path(tmp), 'audit')
            self.assertEqual(raised.exception.status, 403)
            launch.assert_called_once()
            evidence = (Path(tmp) / 'audit_network.json').read_text()
            self.assertIn('ERR_BLOCKED_BY_CLIENT', evidence)
            self.assertNotIn('private-key', evidence)
            self.assertNotIn('?key=', evidence)

    def test_mac_network_diagnostics_distinguish_dns_proxy_and_quota_errors(self):
        from browser_collector import network_problem, network_url
        self.assertIn('hostname on this Mac', network_problem([{'error': 'net::ERR_NAME_NOT_RESOLVED'}])[0])
        self.assertIn('certificate or proxy', network_problem([{'error': 'net::ERR_TUNNEL_CONNECTION_FAILED'}])[0])
        self.assertEqual(network_problem([{'status': 429}])[1], 429)
        self.assertEqual(network_url('https://user:password@example.com/path?key=private#fragment'), 'https://example.com/path')

    def test_report_format_preview_and_blank_excel_download(self):
        client = server.app.test_client()
        page = client.get('/report-format')
        self.assertEqual(page.status_code, 200)
        self.assertIn(b'This is a blank format preview', page.data)
        for name in ('BenchMark', 'TH', 'BL', 'FL', 'SS', 'URL'):
            self.assertIn(f'<summary>{name}</summary>'.encode(), page.data)
        response = client.get('/report-template.xlsx')
        self.assertEqual(response.status_code, 200)
        self.assertIn('attachment', response.headers['Content-Disposition'])
        book = load_workbook(BytesIO(response.data))
        self.assertEqual(book.sheetnames, ['BenchMark', 'TH', 'BL', 'FL', 'SS', 'URL'])
        for name in ('TH', 'BL', 'FL', 'SS'):
            self.assertTrue(all(cell.value is None for row in book[name].iter_rows(min_row=4) for cell in row))
        book.close()
        response.close()
        self.assertIn(b'/report-format', client.get('/').data)

    def test_workbook_web_view_preserves_headers_values_and_escapes_untrusted_content(self):
        from excel_layout import TEMPLATE
        from report_preview import workbook_preview_html
        from openpyxl.comments import Comment
        book = load_workbook(TEMPLATE)
        book['TH']['B4'] = 2.3
        book['TH']['G4'] = 58
        book['TH']['B4'].comment = Comment('<img src=x onerror=alert(1)>', 'test')
        book['URL']['B2'] = '<script>alert(1)</script>'
        book['URL']['B2'].hyperlink = 'javascript:alert(1)'
        page = workbook_preview_html(book)
        self.assertIn('colspan="14"', page)
        self.assertIn('background:#00B0F0', page)
        self.assertIn('background:#FFFF00', page)
        self.assertIn('>2.3</td>', page)
        self.assertIn('>58</td>', page)
        self.assertIn('&lt;script&gt;', page)
        self.assertNotIn('<script>', page)
        self.assertNotIn('<img', page)
        self.assertNotIn('href="javascript:', page)
        book.close()

    def test_rejected_article_is_reported_without_retrying_or_waiting_for_tabs(self):
        context, page, tab = MagicMock(), MagicMock(), MagicMock()
        context.pages = [page]
        tab.is_visible.return_value = False
        evidence = 'PageSpeed Insights\nUnable to resolve https://example.com/a. Try checking the URL for validity.'
        page.locator.side_effect = lambda selector: tab if selector == '#mobile_tab' else page
        page.inner_text.return_value = evidence
        with tempfile.TemporaryDirectory() as tmp, patch('playwright.sync_api.sync_playwright') as start, patch('browser_collector.sys.platform', 'linux'):
            launch = start.return_value.__enter__.return_value.chromium.launch_persistent_context
            launch.return_value = context
            with self.assertRaisesRegex(engine.FetchError, 'Google PageSpeed could not resolve') as raised:
                collect('https://example.com/a', 'mobile', Path(tmp), 'audit')
            self.assertEqual(raised.exception.status, 400)
            launch.assert_called_once()
            tab.wait_for.assert_not_called()
            self.assertEqual((Path(tmp) / 'audit_failure.txt').read_text(), evidence)
            context.close.assert_called_once()

    def test_transient_browser_failure_still_retries(self):
        row = engine.extract_psi(response())
        with patch('browser_collector.collect_once', side_effect=[engine.FetchError('Temporary connection failure'), row]) as attempt:
            self.assertEqual(collect('https://example.com/a', 'mobile', Path('.'), 'audit'), row)
            self.assertEqual(attempt.call_count, 2)

    def test_browser_launch_failure_explains_policy_and_removes_temporary_profile(self):
        from playwright.sync_api import Error as BrowserError

        with tempfile.TemporaryDirectory() as tmp, patch('playwright.sync_api.sync_playwright') as start, patch('browser_collector.sys.platform', 'linux'):
            launch = start.return_value.__enter__.return_value.chromium.launch_persistent_context
            launch.side_effect = BrowserError('Browser closed before connecting')
            with self.assertRaisesRegex(engine.FetchError, 'ask IT.*Google API'):
                collect_once('https://example.com/a', 'mobile', Path(tmp), 'audit')
            options = launch.call_args.kwargs
            self.assertTrue(options['chromium_sandbox'])
            for flag in ('--disable-extensions', '--disable-component-extensions-with-background-pages',
                         '--disable-background-networking', '--disable-component-update'):
                self.assertIn(flag, options['ignore_default_args'])
            profile = Path(options['user_data_dir'])
            self.assertTrue(profile.name.startswith('pagespeed-chrome-'))
            self.assertFalse(profile.exists())

    def test_browser_collection_writes_evidence_and_cleans_up_context(self):
        context = MagicMock()
        page, tab, panel = MagicMock(), MagicMock(), MagicMock()
        context.pages = [page]
        page.url = 'https://pagespeed.web.dev/analysis/example/report?form_factor=mobile'
        tab.get_attribute.return_value = 'true'
        panel.is_visible.return_value = True
        panel.inner_text.return_value = 'selected device measurements'
        page.locator.side_effect = lambda selector: (
            tab if selector == '#mobile_tab' else panel if selector.startswith('[role=') else page
        )
        page.inner_text.return_value = 'full page evidence'
        page.content.return_value = '<html>full page evidence</html>'
        row = engine.extract_psi(response())
        profiles = []

        def launch(**options):
            profile = Path(options['user_data_dir'])
            self.assertTrue(profile.is_dir())
            profiles.append(profile)
            return context

        with tempfile.TemporaryDirectory() as tmp, patch('playwright.sync_api.sync_playwright') as start, patch('browser_collector.sys.platform', 'linux'), \
                patch('browser_collector.parse_visible', return_value=row) as parse:
            start.return_value.__enter__.return_value.chromium.launch_persistent_context.side_effect = launch
            result = collect_once('https://example.com/a', 'mobile', Path(tmp), 'audit')
            self.assertEqual(result, row)
            parse.assert_called_once_with('selected device measurements', 'https://example.com/a',
                                          'mobile', 'URL', page.url)
            self.assertEqual((Path(tmp) / 'audit_page.txt').read_text(), 'full page evidence')
            self.assertEqual(json.loads((Path(tmp) / 'audit_browser.json').read_text()), row)
            context.close.assert_called_once()
        self.assertFalse(profiles[0].exists())

    def test_browser_values_preserve_field_lab_and_assessment(self):
        self.assertEqual(visible_number('First Contentful Paint\n850 ms\n', 'First Contentful Paint', unit='s'), .85)
        self.assertEqual(visible_number('Total Blocking Time\n1.2 s\n', 'Total Blocking Time', unit='ms'), 1200)
        text = '''Core Web Vitals Assessment:
Failed
Largest Contentful Paint (LCP)
2.5 s
Interaction to Next Paint (INP)
200 ms
Cumulative Layout Shift (CLS)
0.1
First Contentful Paint (FCP)
1.6 s
Time to First Byte (TTFB)
0.6 s
Diagnose performance issues
38
Performance
75
Accessibility
54
Best Practices
92
SEO
First Contentful Paint
3.2\u00a0s
Largest Contentful Paint
4.8\u00a0s
Total Blocking Time
2,640\u00a0ms
Cumulative Layout Shift
0.029
Speed Index
15.0\u00a0s
Captured at Sep 8, 2026, 4:55 PM GMT+5:30
Emulated Moto G Power with Lighthouse 13.4.1
'''
        row = parse_visible(text, 'https://example.com/a', 'mobile', 'Origin', 'https://pagespeed.web.dev/analysis/a/id')
        self.assertEqual(row['Performance'], 38)
        self.assertEqual(row['LCP (s)'], 2.5)
        self.assertEqual(row['Lab LCP (s)'], 4.8)
        self.assertEqual(row['Lab TBT (ms)'], 2640)
        self.assertEqual(row['CWV assessment'], 'Failed')
        self.assertEqual(row['Field URL'], 'https://example.com')
        self.assertEqual(row['Lighthouse version'], '13.4.1')
        partial = parse_visible(text.replace('38\nPerformance', 'Error\nPerformance'), 'https://example.com/a', 'mobile', 'Origin', '')
        self.assertEqual(partial['Status'], 'Partial')
        self.assertEqual(partial['LCP (s)'], 2.5)
        self.assertIsNone(partial['Performance'])
        field_only = parse_visible(text.split('Diagnose performance issues')[0], 'https://example.com/a', 'mobile', 'Origin', '')
        self.assertEqual(field_only['Status'], 'Partial')
        self.assertEqual(field_only['INP (ms)'], 200)
        with self.assertRaises(engine.FetchError):
            parse_visible(text, 'https://example.com/a', 'mobile', None, '')

    def test_rss_atom_and_duplicates(self):
        self.assertEqual(len(engine.parse_feed(RSS, "https://example.com/feed")), 2)
        atom = b'<feed xmlns="http://www.w3.org/2005/Atom"><entry><title>A</title><link rel="self" href="/ignore"/><link href="/article"/></entry></feed>'
        self.assertEqual(engine.parse_feed(atom, "https://example.com/feed")[0]["url"], "https://example.com/article")
        with self.assertRaises(ValueError):
            engine.parse_feed(b'<html/>', 'https://example.com')

    def test_metrics_units_and_assessment(self):
        row = engine.extract_psi(response())
        self.assertEqual(row['LCP (s)'], 2.5)
        self.assertEqual(row['CLS'], .1)
        self.assertEqual(row['Performance'], 90)
        self.assertEqual(row['Lab FCP (s)'], 1.25)
        self.assertEqual(row['Lab TBT (ms)'], 500)
        self.assertEqual(row['CWV assessment'], 'Passed')
        row['INP (ms)'] = None
        self.assertEqual(engine.assessment(row), 'Unavailable')
        row['INP (ms)'] = 201
        self.assertEqual(engine.assessment(row), 'Failed')
        self.assertEqual(engine.rating(2.5001, 2.5, 4), 'Needs improvement')
        self.assertIsNone(engine.number('NaN'))

    def test_missing_and_origin(self):
        data = response()
        data['originLoadingExperience'] = data.pop('loadingExperience')
        self.assertEqual(engine.extract_psi(data)['Field source'], 'PSI / Origin')
        data.pop('originLoadingExperience')
        row = engine.extract_psi(data)
        self.assertEqual(row['CWV assessment'], 'Unavailable')
        self.assertIsNone(row['INP (ms)'])
        data['lighthouseResult']['runtimeError'] = {'code': 'FAILED', 'message': 'Navigation failed'}
        self.assertEqual(engine.extract_psi(data)['Status'], 'Error')

    def test_crux_origin_fallback(self):
        data = {'record': {'key': {'origin': 'https://example.com'}, 'metrics': {
            'cumulative_layout_shift': {'percentiles': {'p75': '0.05'}},
            'largest_contentful_paint': {'percentiles': {'p75': 2100}},
            'interaction_to_next_paint': {'percentiles': {'p75': 180}}}}}
        with tempfile.TemporaryDirectory() as tmp, patch('engine.fetch_json', side_effect=[engine.FetchError('No data', 404), data]) as mock:
            row = engine.extract_psi({})
            engine.add_crux(row, 'https://example.com/a', 'mobile', 'test-key', Path(tmp), 'test')
            self.assertEqual(row['Field source'], 'CrUX / Origin')
            self.assertEqual(row['CLS'], .05)
            self.assertEqual(row['CWV assessment'], 'Passed')
            self.assertEqual(mock.call_args.args[1]['formFactor'], 'PHONE')

    def test_complete_pipeline_exports_and_cross_feed_dedup(self):
        config = {'feeds': {'TH': ['https://example.com/feed1', 'https://example.com/feed2']}, 'sample_size': 2, 'strategies': ['mobile', 'desktop'], 'seed': 42}
        with tempfile.TemporaryDirectory() as tmp, patch('engine.fetch', return_value=RSS), patch('engine.fetch_json', return_value=response()):
            folder = engine.run(config, tmp, psi_key='', crux_key='')
            state = json.loads((folder / 'run.json').read_text(encoding='utf-8'))
            self.assertEqual(state['status'], 'Complete')
            self.assertEqual(len(state['samples']), 2)
            self.assertEqual(len(state['results']), 4)
            wb = load_workbook(folder / 'report.xlsx')
            self.assertEqual(wb['Results'].max_row, 5)
            self.assertEqual(wb['TH'].max_row, 5)
            self.assertTrue(wb['URLs']['E2'].value.startswith("'="))
            self.assertEqual(wb['URLs']['A1'].value, 'Date')
            self.assertEqual(wb['URLs']['A2'].value, wb['URL']['A2'].value)
            self.assertEqual(wb['Results'].cell(2, 14).value, 90)
            self.assertEqual(wb.sheetnames[:6], ['BenchMark', 'TH', 'BL', 'FL', 'SS', 'URL'])
            self.assertEqual(wb['TH']['B1'].value, 'Mobile')
            self.assertEqual(wb['TH']['P1'].value, 'Desktop')
            self.assertEqual(wb['TH']['B4'].value, 2.5)
            self.assertEqual(wb['TH']['P4'].value, 2.5)
            self.assertEqual(wb['TH']['G4'].value, 90)
            self.assertEqual(wb['TH']['U4'].value, 90)
            self.assertEqual(wb['URL']['B2'].value, state['samples'][0]['URL'])
            self.assertEqual(wb['URL']['B3'].value, state['samples'][1]['URL'])
            self.assertEqual(wb['TH']['B1'].fill.fgColor.rgb, 'FF00B0F0')
            self.assertEqual(wb['TH']['P1'].fill.fgColor.rgb, 'FFFFFF00')
            self.assertIn('B1:O1', {str(m) for m in wb['TH'].merged_cells.ranges})
            self.assertNotIn('CWV assessment', [c.value for c in wb['Results'][1]])
            content = (folder / 'report.html').read_text(encoding='utf-8')
            self.assertIn('&lt;script&gt;', content)
            self.assertIn('report.xlsx', content)
            self.assertIn('Report by publication', content)
            self.assertIn('<summary>TH</summary>', content)
            self.assertIn('background:#00B0F0', content)
            self.assertNotIn('CWV passed', content)
            self.assertNotIn('Core Web Vitals:', content)
            self.assertNotIn('assessment', content)

    def test_failures_still_export(self):
        config = {'feeds': {'TH': 'https://example.com/feed'}, 'strategies': ['mobile']}
        with tempfile.TemporaryDirectory() as tmp, patch('engine.fetch', side_effect=engine.FetchError('Denied', 403)):
            folder = engine.run(config, tmp, psi_key='', crux_key='')
            state = json.loads((folder / 'run.json').read_text(encoding='utf-8'))
            self.assertEqual(state['status'], 'Failed')
            self.assertTrue((folder / 'report.xlsx').exists())
        with tempfile.TemporaryDirectory() as tmp, patch('engine.fetch', return_value=RSS), patch('engine.fetch_json', side_effect=engine.FetchError('Quota', 429)):
            folder = engine.run(config, tmp, psi_key='', crux_key='')
            state = json.loads((folder / 'run.json').read_text(encoding='utf-8'))
            self.assertEqual(state['status'], 'Failed')
            self.assertEqual(state['results'][0]['Error'], 'Quota')

    def test_browser_failure_continues_to_next_device(self):
        config = {'feeds': {'TH': 'https://example.com/feed'}, 'strategies': ['mobile', 'desktop'], 'collector': 'browser'}
        with tempfile.TemporaryDirectory() as tmp, patch('engine.fetch', return_value=RSS), patch('browser_collector.collect', side_effect=[engine.FetchError('Lighthouse failed'), engine.extract_psi(response())]):
            folder = engine.run(config, tmp, psi_key='', crux_key='')
            state = json.loads((folder / 'run.json').read_text(encoding='utf-8'))
            self.assertEqual(state['status'], 'Partial')
            self.assertEqual(len(state['results']), 2)
            self.assertEqual(state['results'][1]['Status'], 'OK')

    def test_field_only_run_exports_real_values_as_partial(self):
        config = {'feeds': {'TH': 'https://example.com/feed'}, 'strategies': ['mobile'], 'collector': 'browser'}
        row = engine.extract_psi(response())
        row.update({name: None for name in engine.METRICS[5:]})
        row.update(Status='Partial', Error='Lighthouse is overloaded')
        with tempfile.TemporaryDirectory() as tmp, patch('engine.fetch', return_value=RSS), patch('browser_collector.collect', return_value=row):
            folder = engine.run(config, tmp, psi_key='', crux_key='')
            state = json.loads((folder / 'run.json').read_text(encoding='utf-8'))
            self.assertEqual(state['status'], 'Partial')
            wb = load_workbook(folder / 'report.xlsx')
            self.assertEqual(wb['Results'].cell(2, 9).value, 2.5)
            self.assertEqual(wb['Results'].cell(2, 14).value, 'N/A')
            self.assertNotIn('<h1>No measurements collected</h1>', (folder / 'report.html').read_text(encoding='utf-8'))

    def test_ui_and_validation(self):
        client = server.app.test_client()
        self.assertEqual(client.get('/').status_code, 200)
        self.assertEqual(client.post('/runs', json={}).status_code, 403)
        self.assertEqual(client.post('/runs', json={}, headers={'X-Report-Token': server.TOKEN}).status_code, 400)
        self.assertEqual(client.get('/reports/bad/report.xlsx').status_code, 404)
        self.assertEqual(client.get('/', headers={'Host': 'evil.example'}).status_code, 403)
        with self.assertRaises(ValueError):
            engine.validate_config({'feeds': {'Results': 'https://example.com'}})

    def test_resume_preserves_partial_and_completed_results(self):
        from resume_run import resume
        config = {'feeds': {'TH': 'https://example.com/feed'}, 'strategies': ['mobile', 'desktop'], 'collector': 'browser'}
        good = engine.extract_psi(response())
        partial = engine.extract_psi(response())
        partial.update(Status='Partial', Performance=None, Error='Google Lighthouse is overloaded.')
        # Resume should not request already collected data.
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp) / 'run'
            folder.mkdir()
            (folder / 'raw').mkdir()
            sample = {'Publication': 'TH', 'URL': 'https://example.com/a', 'Title': 'A', 'RSS feed': 'https://example.com/feed'}
            rows = [dict(good, **sample, Device='mobile', **{'PageSpeed link': 'https://pagespeed.web.dev/'}), dict(partial, **sample, Device='desktop', **{'PageSpeed link': 'https://pagespeed.web.dev/'})]
            state = {'id': 'test', 'started': engine.now(), 'seed': 1, 'config': config, 'status': 'Partial', 'samples': [sample], 'results': rows, 'errors': []}
            engine.save_json(folder / 'run.json', state)
            with patch('resume_run.collect') as fetch:
                result = resume(folder)
                fetch.assert_not_called()
            self.assertEqual(result['status'], 'Partial')
            self.assertTrue((folder / 'report.xlsx').exists())

    def test_running_server_exports_are_upgraded_once_on_completion(self):
        from refresh_report_layout import refresh
        config = {'feeds': {'TH': 'https://example.com/feed'}, 'strategies': ['mobile']}
        with tempfile.TemporaryDirectory() as tmp, patch('engine.fetch', return_value=RSS), patch('engine.fetch_json', return_value=response()):
            folder = engine.run(config, tmp, psi_key='', crux_key='')
            self.assertEqual(refresh(tmp), [])
            (folder / 'report_layout.json').unlink()
            state = json.loads((folder / 'run.json').read_text(encoding='utf-8'))
            state['status'] = 'Running'
            engine.save_json(folder / 'run.json', state)
            self.assertEqual(refresh(tmp), [])
            state['status'] = 'Complete'
            engine.save_json(folder / 'run.json', state)
            self.assertEqual(refresh(tmp), [folder.name])
            self.assertEqual(refresh(tmp), [])
            wb = load_workbook(folder / 'report.xlsx')
            self.assertEqual(wb['TH']['P1'].value, 'Desktop')
            self.assertNotIn('CWV assessment', [cell.value for cell in wb['Results'][1]])


if __name__ == '__main__':
    unittest.main()
